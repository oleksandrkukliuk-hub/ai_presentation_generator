from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


class Theme:
    """
    Modern Science Education theme.
    """

    # Background
    BACKGROUND = RGBColor(248, 250, 252)

    # Main colors
    NAVY = RGBColor(15, 23, 42)
    BLUE = RGBColor(37, 99, 235)
    CYAN = RGBColor(6, 182, 212)
    PURPLE = RGBColor(124, 58, 237)
    GREEN = RGBColor(16, 185, 129)
    ORANGE = RGBColor(245, 158, 11)
    RED = RGBColor(239, 68, 68)

    # Text
    TEXT = RGBColor(51, 65, 85)
    MUTED = RGBColor(100, 116, 139)

    # UI
    WHITE = RGBColor(255, 255, 255)
    BORDER = RGBColor(226, 232, 240)

    FONT = "Manrope"

    TITLE_SIZE = 40
    SUBTITLE_SIZE = 20
    BODY_SIZE = 22
    SMALL_SIZE = 12

    CARD_TITLE_SIZE = 22
    CARD_BODY_SIZE = 17

    SECTION_TITLE_SIZE = 40
    QUOTE_SIZE = 30



class Presentation_:
    """
    High-level API for creating Modern Science Education presentations.

    Built on top of python-pptx.
    """

    def __init__(
            self,
            path: str = "presentation.pptx",
            width: float = 13.333,
            height: float = 7.5,
            theme=Theme,
    ):
        self.path = Path(path)
        self.theme = theme

        self.width = width
        self.height = height

        self.pptx = Presentation()

        self.pptx.slide_width = Inches(width)
        self.pptx.slide_height = Inches(height)

        self.slide_number = 0

        # Remove default slide
        self._remove_default_slides()

    def save(self, path=None):
        """
        Save presentation to .pptx file.
        """

        output = Path(path) if path else self.path

        output.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        self.pptx.save(output)

        return output

    # =========================================================
    # INTERNAL
    # =========================================================

    def _remove_default_slides(self):
        while len(self.pptx.slides) > 0:
            slide_id = self.pptx.slides._sldIdLst[-1]
            self.pptx.part.drop_rel(slide_id.rId)
            del self.pptx.slides._sldIdLst[-1]

    def _new_slide(self, background=None):
        slide = self.pptx.slides.add_slide(
            self.pptx.slide_layouts[6]
        )

        self.slide_number += 1

        background = background or self.theme.BACKGROUND

        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = background

        self._add_accent_bar(slide)

        return slide

    def _add_accent_bar(self, slide):
        """
        Thin blue vertical bar on the left side.
        """

        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            0,
            0,
            Inches(0.10),
            self.pptx.slide_height,
        )

        shape.fill.solid()
        shape.fill.fore_color.rgb = self.theme.BLUE
        shape.line.fill.background()

    def _add_text(
            self,
            slide,
            text,
            x,
            y,
            width,
            height,
            size=16,
            color=None,
            bold=False,
            align=PP_ALIGN.LEFT,
            valign=MSO_ANCHOR.TOP,
    ):
        box = slide.shapes.add_textbox(
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )

        frame = box.text_frame

        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = valign

        paragraph = frame.paragraphs[0]
        paragraph.alignment = align

        run = paragraph.add_run()
        run.text = str(text)

        run.font.name = self.theme.FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color or self.theme.TEXT

        return box

    def _add_title(self, slide, title, subtitle=None):
        self._add_text(
            slide,
            title,
            0.75,
            0.55,
            11.5,
            0.65,
            size=self.theme.TITLE_SIZE,
            color=self.theme.NAVY,
            bold=True,
        )

        if subtitle:
            self._add_text(
                slide,
                subtitle,
                0.77,
                1.18,
                11.2,
                0.45,
                size=self.theme.SUBTITLE_SIZE,
                color=self.theme.MUTED,
            )

    def _add_footer(self, slide):
        self._add_text(
            slide,
            f"{self.slide_number:02d}",
            12.0,
            7.0,
            0.5,
            0.25,
            size=self.theme.SMALL_SIZE,
            color=self.theme.MUTED,
            align=PP_ALIGN.RIGHT,
        )

    def _add_card(
            self,
            slide,
            x,
            y,
            width,
            height,
            color=None,
    ):
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(width),
            Inches(height),
        )

        card.fill.solid()
        card.fill.fore_color.rgb = color or self.theme.WHITE

        card.line.color.rgb = self.theme.BORDER
        card.line.width = Pt(0.7)

        return card

    # =========================================================
    # PUBLIC API
    # =========================================================

    def add_title_slide(
            self,
            title: str,
            subtitle: str = "",
            label: str = "SCIENCE • EDUCATION",
    ):
        """
        Add modern title slide.
        """

        slide = self._new_slide()

        # Decorative circles
        self._circle(
            slide,
            9.5,
            0.8,
            2.2,
            self.theme.BLUE,
        )

        self._circle(
            slide,
            10.7,
            1.7,
            1.1,
            self.theme.CYAN,
        )

        self._circle(
            slide,
            9.0,
            2.0,
            0.5,
            self.theme.PURPLE,
        )

        self._add_text(
            slide,
            label,
            0.85,
            1.05,
            6,
            0.4,
            size=11,
            color=self.theme.BLUE,
            bold=True,
        )

        self._add_text(
            slide,
            title,
            0.85,
            1.7,
            7.8,
            1.5,
            size=34,
            color=self.theme.NAVY,
            bold=True,
        )

        self._add_text(
            slide,
            subtitle,
            0.88,
            3.4,
            7,
            0.9,
            size=18,
            color=self.theme.TEXT,
        )

        self._add_footer(slide)

        return slide

    def add_content_slide(
            self,
            title: str,
            bullets: list[str],
            subtitle: str | None = None,
    ):
        """
        Standard educational slide.
        """

        slide = self._new_slide()

        self._add_title(
            slide,
            title,
            subtitle,
        )

        y = 1.9

        colors = [
            self.theme.BLUE,
            self.theme.CYAN,
            self.theme.GREEN,
            self.theme.PURPLE,
        ]

        for index, bullet in enumerate(bullets):
            color = colors[index % len(colors)]

            # Bullet
            dot = slide.shapes.add_shape(
                MSO_SHAPE.OVAL,
                Inches(0.9),
                Inches(y + 0.13),
                Inches(0.14),
                Inches(0.14),
            )

            dot.fill.solid()
            dot.fill.fore_color.rgb = color
            dot.line.fill.background()

            # Text
            self._add_text(
                slide,
                bullet,
                1.25,
                y,
                10.5,
                0.7,
                size=self.theme.BODY_SIZE,
                color=self.theme.TEXT,
            )

            y += 0.9

        self._add_footer(slide)

        return slide

    def add_two_column_slide(
            self,
            title: str,
            left_title: str,
            left_content,
            right_title: str,
            right_content,
    ):
        """
        Two-column educational slide.
        """

        slide = self._new_slide()

        self._add_title(
            slide,
            title,
        )

        self._add_card(
            slide,
            0.75,
            1.75,
            5.75,
            4.65,
        )

        self._add_card(
            slide,
            6.8,
            1.75,
            5.75,
            4.65,
        )

        self._add_text(
            slide,
            left_title,
            1.05,
            2.05,
            5,
            0.5,
            size=19,
            color=self.theme.BLUE,
            bold=True,
        )

        self._add_text(
            slide,
            right_title,
            7.1,
            2.05,
            5,
            0.5,
            size=19,
            color=self.theme.PURPLE,
            bold=True,
        )

        self._add_content(
            slide,
            left_content,
            1.05,
            2.75,
            5,
        )

        self._add_content(
            slide,
            right_content,
            7.1,
            2.75,
            5,
        )

        self._add_footer(slide)

        return slide

    def add_comparison_slide(
            self,
            title: str,
            left_title: str,
            left_items: list[str],
            right_title: str,
            right_items: list[str],
    ):
        """
        Comparison slide.
        """

        return self.add_two_column_slide(
            title=title,
            left_title=left_title,
            left_content=left_items,
            right_title=right_title,
            right_content=right_items,
        )

    def add_cards_slide(
            self,
            title: str,
            cards: list[dict],
            columns: int = 3,
    ):
        """
        Card-based educational slide.

        Example:

        cards=[
            {
                "title": "Superposition",
                "text": "...",
                "accent": "blue"
            }
        ]
        """

        slide = self._new_slide()

        self._add_title(
            slide,
            title,
        )

        columns = max(1, min(columns, 4))

        gap = 0.25
        total_width = 11.8

        card_width = (
                             total_width -
                             gap * (columns - 1)
                     ) / columns

        accent_colors = {
            "blue": self.theme.BLUE,
            "cyan": self.theme.CYAN,
            "green": self.theme.GREEN,
            "purple": self.theme.PURPLE,
            "orange": self.theme.ORANGE,
            "red": self.theme.RED,
        }

        for index, card_data in enumerate(cards):
            column = index % columns
            row = index // columns

            x = 0.75 + column * (
                    card_width + gap
            )

            y = 1.75 + row * 2.35

            self._add_card(
                slide,
                x,
                y,
                card_width,
                2.05,
            )

            accent = accent_colors.get(
                card_data.get("accent", "blue"),
                self.theme.BLUE,
            )

            bar = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                Inches(x),
                Inches(y),
                Inches(0.08),
                Inches(2.05),
            )

            bar.fill.solid()
            bar.fill.fore_color.rgb = accent
            bar.line.fill.background()

            self._add_text(
                slide,
                card_data["title"],
                x + 0.3,
                y + 0.3,
                card_width - 0.5,
                0.45,
                size=17,
                color=self.theme.NAVY,
                bold=True,
            )

            self._add_text(
                slide,
                card_data.get("text", ""),
                x + 0.3,
                y + 0.9,
                card_width - 0.5,
                0.85,
                size=12,
                color=self.theme.TEXT,
            )

        self._add_footer(slide)

        return slide

    def add_section_slide(
            self,
            title: str,
            description: str = "",
            number: int | None = None,
    ):
        """
        Dark section divider.
        """

        slide = self._new_slide(
            background=self.theme.NAVY
        )

        self._add_text(
            slide,
            f"{number:02d}" if number else "SECTION",
            0.85,
            1.1,
            2,
            0.4,
            size=12,
            color=self.theme.CYAN,
            bold=True,
        )

        self._add_text(
            slide,
            title,
            0.85,
            2.1,
            7.5,
            1.2,
            size=34,
            color=self.theme.WHITE,
            bold=True,
        )

        self._add_text(
            slide,
            description,
            0.88,
            3.6,
            6.8,
            1,
            size=17,
            color=RGBColor(203, 213, 225),
        )

        self._add_footer(slide)

        return slide

    def add_quote_slide(
            self,
            quote: str,
            author: str = "",
    ):
        """
        Quote / key idea slide.
        """

        slide = self._new_slide()

        self._add_text(
            slide,
            "“",
            0.85,
            1.0,
            1,
            1,
            size=54,
            color=self.theme.BLUE,
            bold=True,
        )

        self._add_text(
            slide,
            quote,
            1.35,
            1.75,
            10.3,
            2.5,
            size=27,
            color=self.theme.NAVY,
            bold=True,
            valign=MSO_ANCHOR.MIDDLE,
        )

        if author:
            self._add_text(
                slide,
                f"— {author}",
                1.4,
                4.65,
                7,
                0.45,
                size=14,
                color=self.theme.MUTED,
            )

        self._add_footer(slide)

        return slide

    def add_image_slide(
            self,
            title: str,
            image_path: str,
            caption: str = "",
    ):
        """
        Large image slide.
        """

        slide = self._new_slide()

        self._add_title(
            slide,
            title,
        )

        slide.shapes.add_picture(
            image_path,
            Inches(0.95),
            Inches(1.65),
            width=Inches(11.4),
            height=Inches(4.7),
        )

        if caption:
            self._add_text(
                slide,
                caption,
                1,
                6.4,
                11.2,
                0.4,
                size=10,
                color=self.theme.MUTED,
                align=PP_ALIGN.CENTER,
            )

        self._add_footer(slide)

        return slide

    # =========================================================
    # HELPERS
    # =========================================================

    def _circle(
            self,
            slide,
            x,
            y,
            size,
            color,
    ):
        circle = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            Inches(x),
            Inches(y),
            Inches(size),
            Inches(size),
        )

        circle.fill.background()
        circle.line.color.rgb = color
        circle.line.width = Pt(2)

        return circle

    def _add_content(
            self,
            slide,
            content,
            x,
            y,
            width,
    ):
        if isinstance(content, (list, tuple)):

            for index, item in enumerate(content):
                self._add_text(
                    slide,
                    f"• {item}",
                    x,
                    y + index * 0.72,
                    width,
                    0.55,
                    size=self.theme.BODY_SIZE,
                    color=self.theme.TEXT,
                )

        else:

            self._add_text(
                slide,
                content,
                x,
                y,
                width,
                2.8,
                size=self.theme.BODY_SIZE,
                color=self.theme.TEXT,
            )


if __name__ == "__main__":
    p = Presentation_(
        "quantum_computers.pptx"
    )

    p.add_title_slide(
        title="Квантові комп'ютери",
        subtitle="Принцип роботи та перспективи",
    )

    p.add_content_slide(
        title="Що таке кубіт?",
        bullets=[
            "Кубіт — базова одиниця квантової інформації",
            "Може перебувати у суперпозиції станів",
            "Використовує квантову інтерференцію",
            "Вимірювання визначає результат",
        ],
    )

    p.add_comparison_slide(
        title="Класичний біт vs кубіт",
        left_title="Класичний біт",
        left_items=[
            "0 або 1",
            "Один визначений стан",
            "Класичні операції",
        ],
        right_title="Кубіт",
        right_items=[
            "0, 1 або суперпозиція",
            "Квантовий стан",
            "Квантові операції",
        ],
    )

    p.add_cards_slide(
        title="Основні принципи",
        cards=[
            {
                "title": "Суперпозиція",
                "text": "Кубіт може перебувати в комбінації станів.",
                "accent": "blue",
            },
            {
                "title": "Заплутаність",
                "text": "Квантові системи можуть мати корельовані стани.",
                "accent": "purple",
            },
            {
                "title": "Інтерференція",
                "text": "Квантові амплітуди можуть підсилюватися або гаситися.",
                "accent": "cyan",
            },
        ],
    )

    p.add_section_slide(
        title="Квантові алгоритми",
        description="Як квантові властивості використовуються для обчислень",
        number=2,
    )

    p.save()
